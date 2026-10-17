# -*- coding: utf-8 -*-
"""Собирает .docx для RAG-индекса из переведённых кодексов (storage/app/official/derived/translated/*.json).

Запуск из корня проекта:  python python/tools/build_translated_codes.py
Каждая статья оформляется как «Статья N. Название» (так её режет индексатор), главы — «ГЛАВА N. Название».
После сборки: перезапустить сервер (индекс пересоберётся) и выполнить python python/tools/build_embeddings.py.
"""
import json
import os
import sys

from docx import Document

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", ".."))
SRC = os.path.join(ROOT, "storage", "app", "official", "derived", "translated")
OUT = os.path.join(ROOT, "storage", "app", "rag")

# slug -> имя файла в индексе (должно совпадать с CHAT_RAG_SOURCE_NAMES в app.py)
FILES = {
    "konstitutsiya": "Конституция РТ (машинный перевод).docx",
    "zemelny": "Земельный кодекс (машинный перевод).docx",
    "uik": "Уголовно-исполнительный кодекс (машинный перевод).docx",
    "tamozhenny": "Таможенный кодекс (машинный перевод).docx",
}


def build(slug, name):
    path = os.path.join(SRC, slug + ".json")
    if not os.path.isfile(path):
        print("пропуск (нет перевода):", slug)
        return
    data = json.load(open(path, encoding="utf-8"))
    articles = sorted(data["articles"], key=lambda a: a["pos"])
    chapters = sorted(data.get("chapters") or [], key=lambda c: c["pos"])
    document = Document()
    ci = 0
    for article in articles:
        while ci < len(chapters) and chapters[ci]["pos"] < article["pos"]:
            chapter = chapters[ci]
            document.add_paragraph("ГЛАВА %s. %s" % (chapter["n"], chapter["title_ru"]))
            ci += 1
        head = "Статья %s." % article["n"]
        if article.get("title_ru"):
            head += " " + article["title_ru"]
        document.add_paragraph(head)
        for line in article["text_ru"].split("\n"):
            if line.strip():
                document.add_paragraph(line.strip())
    document.save(os.path.join(OUT, name))
    print("%-14s статей %4d -> %s" % (slug, len(articles), name))


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    for slug, name in FILES.items():
        build(slug, name)
