# -*- coding: utf-8 -*-
"""Справочник размеров штрафов КоАП РТ для RAG-индекса (storage/app/rag/КоАП РТ — размеры штрафов.docx).

Источник: storage/app/official/derived/koap_fines_2026.json — штрафы, разобранные из официального текста КоАП
(base.mmk.tj, sanadID=4) и пересчитанные в сомони по показателю для расчётов (ст. 24 Закона о бюджете на 2026 год = 78).
Каждая статья оформлена как «Статья N. Название» с короткой справкой: так её режет индексатор (legal_rag), и чат
находит и цитирует её как обычную статью кодекса.

Запуск из корня проекта:  python python/tools/build_koap_docx.py
Показатель меняется каждый год: обновите koap_fines_2026.json (parse + пересчёт) и пересоберите документ.
"""
import json
import os
import sys

from docx import Document

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", ".."))
SRC = os.path.join(ROOT, "storage", "app", "official", "derived", "koap_fines_2026.json")
OUT = os.path.join(ROOT, "storage", "app", "rag", "КоАП РТ — размеры штрафов.docx")

EDITION_NOTE = "Редакция КоАП в базе — до 2019 года: если статью меняли позже, размер может быть другим."


def number(value):
    value = float(value)
    text = f"{value:,.0f}" if value.is_integer() else f"{value:,.1f}"
    return text.replace(",", " ")


def phrase(fine, pr):
    lo, hi = fine["min_pr"], fine["max_pr"]
    if lo == hi:
        return f"{number(lo)} показателей для расчётов, то есть {number(fine['min_somoni'])} сомони"
    return (f"от {number(lo)} до {number(hi)} показателей для расчётов, то есть от {number(fine['min_somoni'])} "
            f"до {number(fine['max_somoni'])} сомони")


def main():
    data = json.load(open(SRC, encoding="utf-8"))
    pr = data["pr_somoni"]
    titles = data["titles_ru"]
    descriptions = data.get("descriptions_ru", {})
    document = Document()
    document.add_paragraph("Кодекс Республики Таджикистан об административных правонарушениях — размеры административных штрафов")
    document.add_paragraph(
        f"Справочник размеров штрафов по статьям Особенной части. Один показатель для расчётов равен {pr} сомони "
        f"(Закон РТ о государственном бюджете на 2026 год, статья 24, с 1 января 2026 года)."
    )
    for index, article in enumerate(data["articles"]):
        title = titles.get(f"{article['article']}|{index}") or article["title_tj"]
        document.add_paragraph(f"Статья {article['article']}. {title}")
        description = descriptions.get(f"{article['article']}|{index}")
        has_cases = all(fine.get("cond_ru") for fine in article["fines"])
        if description and not has_cases:
            document.add_paragraph(f"Что считается правонарушением (краткое изложение текста статьи, машинный перевод): {description}")
        # Порядок важен: в контекст чата на статью уходит около 800 символов, поэтому случаи и суммы идут первыми,
        # а слова «как это называют люди» (нужны только поиску) — в конце. Не начинать абзац со слова «Статья»:
        # индексатор примет его за заголовок новой статьи.
        document.add_paragraph(f"Штрафы по статье {article['article']} КоАП РТ (1 показатель для расчётов = {pr} сомони в 2026 году):")
        current_case = None
        keywords = []
        # За первое нарушение по части статей назначается только предупреждение, штраф — за повторное: об этом нельзя молчать.
        for warning in article.get("warnings") or []:
            if warning.get("cond_ru"):
                document.add_paragraph(f"Случай: {warning['cond_ru']}.")
                document.add_paragraph("— За первое такое нарушение назначается предупреждение, штраф не назначается.")
                if warning.get("kw_ru"):
                    keywords.append(warning["kw_ru"])
        for fine in article["fines"]:
            case = fine.get("cond_ru")
            if case and case != current_case:
                document.add_paragraph(f"Случай: {case}.")
                current_case = case
                if fine.get("kw_ru"):
                    keywords.append(fine["kw_ru"])
            tail = " (вместо штрафа возможно предупреждение)" if fine.get("or_warning") else ""
            document.add_paragraph(f"— {fine['subject']}: {phrase(fine, pr)}{tail}.")
        document.add_paragraph(EDITION_NOTE)
        if keywords:
            document.add_paragraph("Как это называют обычными словами: " + "; ".join(keywords)[:900] + ".")
    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    document.save(OUT)
    print("статей:", len(data["articles"]), "->", OUT)


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    main()
